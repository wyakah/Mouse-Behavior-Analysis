import hashlib
import json
from pathlib import Path
import pytest
from flask import Flask
from desktop.service import install_workspace, secure_app, mark_interrupted, active_jobs
from threechamber.workbooks import export_stereotypy, export_three_chamber
from openpyxl import load_workbook

def payload(tmp):
    bundle=tmp/'bundle';(bundle/'payload/static').mkdir(parents=True)
    source=bundle/'payload/static/app.js';source.write_text('test')
    (bundle/'manifest.json').write_text(json.dumps(dict(version='0.1.0',program_files={'static/app.js':hashlib.sha256(b'test').hexdigest()},model_files={})))
    return bundle

def test_install_preserves_user_files_and_updates_owned_files(tmp_path):
    bundle=payload(tmp_path);workspace=tmp_path/'workspace'
    install_workspace(bundle,workspace)
    (workspace/'videos/experiment.mov').write_bytes(b'keep')
    (workspace/'configs/mouse.json').write_text('{}')
    install_workspace(bundle,workspace)
    assert (workspace/'videos/experiment.mov').read_bytes()==b'keep'
    assert (workspace/'configs/mouse.json').read_text()=='{}'
    assert (workspace/'.dlc-env').is_symlink()
    assert (workspace/'static/app.js').read_text()=='test'

def test_install_refuses_tampered_program_and_traversal(tmp_path):
    bundle=payload(tmp_path)
    (bundle/'payload/static/app.js').write_text('damaged')
    with pytest.raises(ValueError,match='damaged'):install_workspace(bundle,tmp_path/'w')
    (bundle/'manifest.json').write_text(json.dumps(dict(version='0.1.0',program_files={'../bad':'x'},model_files={})))
    with pytest.raises(ValueError,match='Invalid'):install_workspace(bundle,tmp_path/'w')

def test_private_loopback_session_requires_bootstrap_cookie():
    app=Flask(__name__);app.config['DESKTOP_WORKSPACE']='/test';secure_app(app,'test-token')
    @app.get('/')
    def index():return 'ready'
    c=app.test_client()
    assert c.get('/').status_code==403
    assert c.get('/desktop/connect?token=wrong').status_code==403
    response=c.get('/desktop/connect?token=test-token')
    assert response.status_code==302 and response.location=='/'
    assert 'HttpOnly' in response.headers['Set-Cookie']
    assert c.get('/').data==b'ready'
    assert c.get('/api/desktop').json['offline'] is True

def test_recovery_preserves_completed_entries(tmp_path):
    p=tmp_path/'outputs/stereo-test/batch.json';p.parent.mkdir(parents=True)
    p.write_text(json.dumps(dict(status='running',entries=[dict(status='complete',summary=[1]),dict(status='running')])))
    assert active_jobs(tmp_path)
    mark_interrupted(tmp_path)
    b=json.loads(p.read_text());assert not active_jobs(tmp_path)
    assert b['entries'][0]['summary']==[1] and b['entries'][0]['status']=='complete'
    assert b['entries'][1]['status']=='failed'

def test_portable_excel_exports_keep_values_and_prevent_formulas(tmp_path):
    dest=tmp_path/'scores.xlsx';export_stereotypy(dict(headers=['Mouse','Seconds'],rows=[['=HYPERLINK("bad")',12.5],['001',None]]),dest)
    ws=load_workbook(dest).active
    assert ws['A2'].data_type=='s' and ws['A2'].value.startswith('=')
    assert ws['B2'].value==12.5 and ws['B3'].value is None and ws['A3'].value=='001'
    assert ws.freeze_panes=='A2'
    batch=dict(entries=[dict(id='001',video='videos/a.mp4',status='complete',config={},summary={'target_side':'left','nose_scoreable_fraction':1,'left_nose_seconds':9,'right_nose_seconds':4},bouts=[])],statistics_report=dict(groups=[dict(n=2,mean=3)],settings={},versions={},notes=[],sources=[],test_count=0))
    export_three_chamber(batch,dest);book=load_workbook(dest)
    assert set(['Results','Setup','Bouts','Groups','Statistics notes'])<=set(book.sheetnames)
    row=dict(zip([c.value for c in book['Results'][1]],[c.value for c in book['Results'][2]]))
    assert row['target_nose_seconds']==9 and row['other_nose_seconds']==4
