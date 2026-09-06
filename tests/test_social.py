import numpy as np
import pytest
from threechamber.social import stranger_metrics
from threechamber.core import score
from threechamber.batch import validate_batch
from threechamber.statistics import metric_value
from test_batch import batch_fixture,circle_cfg
from test_analysis import tracks

def summary(**extra):
    return dict(target_side='right',left_nose_seconds=8.,right_nose_seconds=12.,left_chamber_seconds=10.,center_chamber_seconds=20.,right_chamber_seconds=30.,unknown_chamber_seconds=40.,outside_chamber_seconds=20.,nose_scoreable_fraction=.8,**extra)

def test_stranger_denominator_is_all_chambers_not_cup_time_or_recording():
    s=summary();r=stranger_metrics(s)
    assert r['stranger_interaction_seconds']==12
    assert r['total_chamber_seconds']==60
    assert r['stranger_interaction_percent']==20
    s['stranger_side']='left'
    assert stranger_metrics(s)['stranger_interaction_percent']==pytest.approx(100*8/60)

def test_missing_data_stays_blank_but_observed_zero_is_zero():
    s=summary();s['right_nose_seconds']=0
    assert stranger_metrics(s)['stranger_interaction_percent']==0
    for edit in [dict(target_side='unspecified'),dict(nose_scoreable_fraction=0),dict(right_nose_seconds=None),dict(left_chamber_seconds=0,center_chamber_seconds=0,right_chamber_seconds=0)]:
        assert stranger_metrics({**s,**edit})['stranger_interaction_percent'] is None

def test_timestamp_weighted_score_exports_stranger_side_and_percent():
    cfg={**circle_cfg(),'stranger_side':'right'}
    rows,s=score(tracks([[50,100],[250,100],[250,100],[150,100]]),np.array([0.,1.,3.,6.]),np.array([1.,2.,3.,4.]),cfg)
    assert s['stranger_side']=='right' and s['target_side']=='right'
    assert s['total_chamber_seconds']==10
    assert s['stranger_interaction_seconds']==5 and s['stranger_interaction_percent']==50
    assert rows.stranger_interaction.tolist()==[False,True,True,False]
    assert metric_value(dict(status='complete',summary=s),'stranger_interaction_percent')==(50,None)

def test_batch_requires_position_and_persists_metadata_into_scoring_config(tmp_path):
    draft=batch_fixture(tmp_path)
    for e in draft['entries']:e['config'].pop('target_side')
    with pytest.raises(ValueError,match='stranger mouse position'):validate_batch(tmp_path,draft)
    for e in draft['entries']:e['stranger_side']='right'
    out=validate_batch(tmp_path,draft)
    assert all(e['config']['stranger_side']==e['config']['target_side']=='right' for e in out['entries'])
    draft['entries'][0]['stranger_side']='middle'
    with pytest.raises(ValueError,match='left or right'):validate_batch(tmp_path,draft)


def test_excel_chamber_roles_remain_correct_for_mixed_stranger_sides(tmp_path):
    from openpyxl import load_workbook
    from threechamber.workbooks import export_three_chamber
    entries=[dict(id=side,video=side+'.mp4',status='complete',summary=dict(summary(),stranger_side=side)) for side in ('left','right')]
    file=tmp_path/'roles.xlsx';export_three_chamber({'entries':entries},file)
    rows=list(load_workbook(file)['Results'].values)
    data=[dict(zip(rows[0],r)) for r in rows[1:]]
    assert data[0]['Left chamber']=='Left (Stranger)'
    assert data[1]['Left chamber']=='Left (Object)'
    assert all(r['Left Chamber Time (s)']==10 for r in data)
    export_three_chamber({'entries':entries[:1]},file)
    headers=next(load_workbook(file)['Results'].values)
    assert 'Left Chamber (Stranger) Time (s)' in headers and 'Right Chamber (Object) Time (s)' in headers
    assert 'Chamber Stranger Interaction %' in headers
    assert 'Zone Stranger Interaction %' in headers
    assert 'Center Zone Time (s)' in headers


def test_distinct_chamber_and_zone_metrics_and_missing_nose():
    s={**summary(), 'left_chamber_seconds':200.,'center_chamber_seconds':100.,'right_chamber_seconds':100.,'right_nose_seconds':20.}
    r=stranger_metrics(s)
    assert r['chamber_stranger_interaction_percent']==25
    assert r['zone_stranger_interaction_percent']==5
    assert r['right_zone_seconds']==20
    assert r['center_zone_seconds'] is None
    r=stranger_metrics({**s,'nose_scoreable_fraction':0})
    assert r['chamber_stranger_interaction_percent']==25
    assert r['zone_stranger_interaction_percent'] is None
    assert r['left_zone_seconds'] is None
    r=stranger_metrics({**s,'stranger_side':'left'})
    assert r['chamber_stranger_interaction_percent']==50
    assert r['zone_stranger_interaction_percent']==2
